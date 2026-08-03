import pandas as pd
from pathlib import Path
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path("Scrape PHEI")
OUTPUT_DIR = Path("output_eom")
OUTPUT_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = OUTPUT_DIR / "Selected_EoM_Yield_Curve.xlsx"

month_map = {
    "Januari": 1,
    "Februari": 2,
    "Maret": 3,
    "April": 4,
    "Mei": 5,
    "Juni": 6,
    "Juli": 7,
    "Agustus": 8,
    "September": 9,
    "Oktober": 10,
    "November": 11,
    "Desember": 12,
}

def parse_indonesian_date(date_text):
    """
    Expected file name pattern:
    Yield-Curve-30-Desember-2024.xlsx
    or
    Yield-Curve- 30-Desember-2024.xlsx
    """
    date_text = date_text.strip()
    parts = date_text.split("-")

    if len(parts) != 3:
        raise ValueError(f"Unexpected date format: {date_text}")

    day = int(parts[0].strip())
    month_name = parts[1].strip()
    year = int(parts[2].strip())

    month = month_map[month_name]

    return pd.Timestamp(year=year, month=month, day=day)

def find_selected_files():
    records = []

    for file in BASE_DIR.rglob("Yield-Curve-*.xlsx"):
        match = re.search(r"Yield-Curve-\s*(.+)\.xlsx$", file.name)

        if not match:
            continue

        date_text = match.group(1).strip()

        try:
            file_date = parse_indonesian_date(date_text)
        except Exception as e:
            print(f"Skip file because date cannot be parsed: {file.name}. Error: {e}")
            continue

        keep_2026_monthly = file_date.year == 2026
        keep_dec_2024_2025 = file_date.year in [2024, 2025] and file_date.month == 12

        if keep_2026_monthly or keep_dec_2024_2025:
            records.append({
                "file_path": file,
                "date": file_date,
                "year": file_date.year,
                "month": file_date.month,
                "file_name": file.name
            })

    files_df = pd.DataFrame(records)

    if files_df.empty:
        raise ValueError("No matching Yield-Curve files found for 2026 monthly or Dec 2024/Dec 2025.")

    selected_files = (
        files_df
        .sort_values("date")
        .groupby(["year", "month"], as_index=False)
        .tail(1)
        .sort_values("date")
        .reset_index(drop=True)
    )

    return selected_files

def read_and_compile_long(selected_files):
    compiled_data = []

    for _, row in selected_files.iterrows():
        source_file = row["file_path"]
        source_date = row["date"]

        df = pd.read_excel(source_file, index_col=0)
        df = df.reset_index()

        first_col = df.columns[0]
        df = df.rename(columns={first_col: "Tenor Year"})

        df.insert(0, "Date", source_date)
        df.insert(1, "Year", row["year"])
        df.insert(2, "Month", row["month"])

        if row["year"] in [2024, 2025] and row["month"] == 12:
            df.insert(3, "Period Type", "Year End")
        else:
            df.insert(3, "Period Type", "Month End")

        df["Source File"] = row["file_name"]

        compiled_data.append(df)

    final_df = pd.concat(compiled_data, ignore_index=True)

    return final_df

def make_spot_wide(final_df):
    spot_wide = (
        final_df
        .pivot_table(
            index="Tenor Year",
            columns="Date",
            values="Spot-Rate",
            aggfunc="first"
        )
        .sort_index()
    )

    spot_wide = spot_wide.reindex(sorted(spot_wide.columns), axis=1)

    return spot_wide

def make_yield_wide(final_df):
    yield_wide = (
        final_df
        .pivot_table(
            index="Tenor Year",
            columns="Date",
            values="IBPA Yield",
            aggfunc="first"
        )
        .sort_index()
    )

    yield_wide = yield_wide.reindex(sorted(yield_wide.columns), axis=1)

    return yield_wide

def make_forward_wide(spot_wide):
    """
    Forward rate formula suggestion:
    f_t = ((1 + s_t)^t / (1 + s_{t-1})^(t-1)) - 1

    For t = 1, forward rate is set equal to the 1-year spot rate.
    This section uses integer tenors only.
    """
    integer_tenors = []

    for tenor in spot_wide.index:
        try:
            tenor_float = float(tenor)
            if tenor_float >= 1 and abs(tenor_float - round(tenor_float)) < 1e-9:
                integer_tenors.append(int(round(tenor_float)))
        except Exception:
            continue

    integer_tenors = sorted(set(integer_tenors))

    forward_wide = pd.DataFrame(index=integer_tenors, columns=spot_wide.columns, dtype=float)

    for date_col in spot_wide.columns:
        spot_series = spot_wide[date_col]

        spot_by_tenor = {}
        for tenor in spot_series.index:
            try:
                tenor_float = float(tenor)
                if abs(tenor_float - round(tenor_float)) < 1e-9:
                    spot_by_tenor[int(round(tenor_float))] = spot_series.loc[tenor]
            except Exception:
                continue

        for t in integer_tenors:
            s_t = spot_by_tenor.get(t)

            if pd.isna(s_t):
                forward_wide.loc[t, date_col] = None
                continue

            if t == 1:
                forward_wide.loc[t, date_col] = s_t
            else:
                s_prev = spot_by_tenor.get(t - 1)

                if pd.isna(s_prev) or s_prev is None:
                    forward_wide.loc[t, date_col] = None
                else:
                    forward_wide.loc[t, date_col] = ((1 + s_t) ** t / (1 + s_prev) ** (t - 1)) - 1

    return forward_wide

def write_basic_sheet(writer):
    how_to_use = pd.DataFrame({
        "No": [1, 2, 3],
        "Procedure": [
            "Use this workbook to compare selected IDR spot rates and forward rates.",
            "The IDR data is compiled from selected PHEI yield curve scrapes.",
            "For this automated output, valuation dates are selected as Dec 2024, Dec 2025, and each available month-end in 2026."
        ]
    })

    source_data = pd.DataFrame({
        "Sheet Name": ["IDR", "USD"],
        "Details": [
            "IDR data is compiled from automated PHEI yield curve scraper outputs in this repository.",
            "USD data is not populated by this automation because the current scraper source is IDR/PHEI only."
        ]
    })

    how_to_use.to_excel(writer, sheet_name="How To Use", index=False)
    source_data.to_excel(writer, sheet_name="Source Data", index=False)

def write_idr_sheet(writer, spot_wide, forward_wide, yield_wide):
    workbook = writer.book
    ws = workbook.create_sheet("IDR")

    date_cols = list(spot_wide.columns)

    # Main title
    ws["A1"] = "IDR"
    ws["A1"].font = Font(bold=True, size=14)

    # Spot Rate section
    ws["A3"] = "Spot Rate (IDR) ---- Spot rate of t=3 means the interest rate is effective from t=0 to t=3"
    ws["A3"].font = Font(bold=True)

    ws["A5"] = "Time Period"
    ws["A6"] = "Rate Type"
    ws["A7"] = "Tenor Year"

    for col_idx, date_value in enumerate(date_cols, start=2):
        cell = ws.cell(row=5, column=col_idx)
        cell.value = date_value
        cell.number_format = "mm/dd/yyyy"

        ws.cell(row=6, column=col_idx).value = "Spot"

    start_row = 8

    for row_idx, tenor in enumerate(spot_wide.index, start=start_row):
        ws.cell(row=row_idx, column=1).value = tenor

        for col_idx, date_value in enumerate(date_cols, start=2):
            ws.cell(row=row_idx, column=col_idx).value = spot_wide.loc[tenor, date_value]
            ws.cell(row=row_idx, column=col_idx).number_format = "0.0000%"

    spot_end_row = start_row + len(spot_wide.index) - 1

    # Forward Rate section
    forward_title_row = spot_end_row + 4
    ws.cell(row=forward_title_row, column=1).value = "Forward Rate (IDR) ---- Forward rate of t=3 means the interest rate is effective from t=2 to t=3"
    ws.cell(row=forward_title_row, column=1).font = Font(bold=True)

    forward_header_row = forward_title_row + 2
    ws.cell(row=forward_header_row, column=1).value = "Time Period"
    ws.cell(row=forward_header_row + 1, column=1).value = "Rate Type"
    ws.cell(row=forward_header_row + 2, column=1).value = "Tenor Year"

    for col_idx, date_value in enumerate(date_cols, start=2):
        cell = ws.cell(row=forward_header_row, column=col_idx)
        cell.value = date_value
        cell.number_format = "mm/dd/yyyy"

        ws.cell(row=forward_header_row + 1, column=col_idx).value = "Forward"

    forward_start_row = forward_header_row + 3

    for row_idx, tenor in enumerate(forward_wide.index, start=forward_start_row):
        ws.cell(row=row_idx, column=1).value = tenor

        for col_idx, date_value in enumerate(date_cols, start=2):
            ws.cell(row=row_idx, column=col_idx).value = forward_wide.loc[tenor, date_value]
            ws.cell(row=row_idx, column=col_idx).number_format = "0.0000%"

    # Optional IBPA Yield section for audit
    yield_title_row = forward_start_row + len(forward_wide.index) + 4
    ws.cell(row=yield_title_row, column=1).value = "IBPA Yield (IDR) ---- Original yield values before spot-rate conversion"
    ws.cell(row=yield_title_row, column=1).font = Font(bold=True)

    yield_header_row = yield_title_row + 2
    ws.cell(row=yield_header_row, column=1).value = "Time Period"
    ws.cell(row=yield_header_row + 1, column=1).value = "Rate Type"
    ws.cell(row=yield_header_row + 2, column=1).value = "Tenor Year"

    for col_idx, date_value in enumerate(date_cols, start=2):
        cell = ws.cell(row=yield_header_row, column=col_idx)
        cell.value = date_value
        cell.number_format = "mm/dd/yyyy"

        ws.cell(row=yield_header_row + 1, column=col_idx).value = "Yield"

    yield_start_row = yield_header_row + 3

    for row_idx, tenor in enumerate(yield_wide.index, start=yield_start_row):
        ws.cell(row=row_idx, column=1).value = tenor

        for col_idx, date_value in enumerate(date_cols, start=2):
            ws.cell(row=row_idx, column=col_idx).value = yield_wide.loc[tenor, date_value]
            ws.cell(row=row_idx, column=col_idx).number_format = "0.0000%"

    # Styling
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    thin_bottom = Border(bottom=Side(style="thin", color="808080"))

    section_rows = [3, forward_title_row, yield_title_row]
    header_rows = [5, 6, 7, forward_header_row, forward_header_row + 1, forward_header_row + 2, yield_header_row, yield_header_row + 1, yield_header_row + 2]

    max_col = len(date_cols) + 1

    for row in section_rows:
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = dark_fill
            cell.font = white_font

    for row in header_rows:
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = light_fill
            cell.font = bold_font
            cell.border = thin_bottom

    ws.freeze_panes = "B8"

    ws.column_dimensions["A"].width = 16

    for col in range(2, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 13

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")

def write_usd_placeholder(writer):
    workbook = writer.book
    ws = workbook.create_sheet("USD")
    ws["A1"] = "USD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "USD data is not populated by this automation."
    ws["A4"] = "The current automation compiles IDR data from PHEI yield curve scraper outputs only."

def format_support_sheets(path):
    wb = load_workbook(path)

    for ws in wb.worksheets:
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

        for row in range(1, min(ws.max_row, 10) + 1):
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).alignment = Alignment(vertical="center")

        if ws.max_row >= 1:
            for cell in wscell.font = Font(bold=True)

    wb.save(path)

def main():
    selected_files = find_selected_files()
    final_df = read_and_compile_long(selected_files)

    spot_wide = make_spot_wide(final_df)
    yield_wide = make_yield_wide(final_df)
    forward_wide = make_forward_wide(spot_wide)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        write_basic_sheet(writer)
        write_idr_sheet(writer, spot_wide, forward_wide, yield_wide)
        write_usd_placeholder(writer)

        final_df.to_excel(writer, sheet_name="EoM_Selected_Long", index=False)
        selected_files.to_excel(writer, sheet_name="Selected_Source_Files", index=False)

    format_support_sheets(OUTPUT_FILE)

    print(f"Done. Output saved to: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()